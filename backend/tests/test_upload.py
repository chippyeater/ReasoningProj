import asyncio
import io
import json
import unittest
from unittest.mock import patch

from support import IsolatedCaseTest, docx_bytes
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from starlette.datastructures import UploadFile
from app import case_service, case_store, file_parsers
from app.case_service import append_files_to_case, build_extraction_input
from app.errors import AppError
from app.schemas import FileType


class TestUploads(IsolatedCaseTest):
    def test_upload_stores_exact_original_and_full_utf8_text(self):
        text = "第一段😀\r\n" + "完整正文" * 200 + "\r\n末段"
        raw = text.encode("utf-8")
        case = self.upload(raw=raw)
        file = case["files"][0]
        directory = self.data_dir / case["case_id"]
        self.assertEqual((directory / file["original_path"]).read_bytes(), raw)
        cached = (directory / file["parsed_text_path"]).read_bytes()
        self.assertEqual(cached.decode("utf-8"), text.replace("\r\n", "\n"))
        self.assertNotIn(b"\r", cached)
        self.assertEqual(file["file_size"], len(raw))
        self.assertEqual(file["parse_status"], "parsed")
        self.assertEqual(json.loads((directory / "case.json").read_text(encoding="utf-8")), case)
        self.assertNotIn("完整正文", (directory / "case.json").read_text(encoding="utf-8"))
        run = case["stage_runs"][0]
        self.assertEqual((run["stage"], run["status"]), ("parse", "success"))
        for path in (run["input_snapshot_path"], run["output_snapshot_path"]):
            self.assertTrue((directory / path).is_file())

    def test_append_actually_persists_without_overwriting_case(self):
        before = self.upload()
        first_path = self.data_dir / before["case_id"] / before["files"][0]["parsed_text_path"]
        first_text = first_path.read_bytes()
        response = self.client.post(f"/api/cases/{before['case_id']}/files",
                                    data={"source_kinds": json.dumps({"second.docx": "testimony"})},
                                    files={"files": ("second.docx", docx_bytes("新增段落"))})
        self.assertEqual(response.status_code, 200, response.text)
        after = response.json()["case"]
        self.assertEqual(len(after["files"]), 2)
        self.assertEqual(after["files"][0], before["files"][0])
        self.assertEqual(first_path.read_bytes(), first_text)
        self.assertEqual(after["created_at"], before["created_at"])
        reopened = self.client.get(f"/api/cases/{before['case_id']}")
        self.assertEqual(reopened.json()["case"], after)

    def test_multiple_files_share_one_case_and_report_each_result(self):
        response = self.client.post("/api/cases", data={"case_title": "混合输入", "source_kinds": json.dumps({
            "valid.txt": "document", "broken.docx": "testimony", "scan.png": "record"})}, files=[
            ("files", ("valid.txt", b"valid")), ("files", ("broken.docx", b"not-docx")),
            ("files", ("scan.png", b"no-ocr"))])
        self.assertEqual(response.status_code, 201, response.text)
        case = response.json()["case"]
        self.assertEqual([f["parse_status"] for f in case["files"]], ["parsed", "failed", "unsupported"])
        self.assertEqual([r["status"] for r in case["stage_runs"]], ["success", "failed", "failed"])
        self.assertEqual(self.client.get(f"/api/cases/{case['case_id']}/extraction-input").status_code, 422)
        selected = self.client.get(f"/api/cases/{case['case_id']}/extraction-input",
                                   params={"file_id": case["files"][0]["file_id"]})
        self.assertEqual(selected.status_code, 200, selected.text)
        self.assertEqual(selected.json()["files"][0]["text"], "valid")

    def test_empty_corrupt_invalid_utf8_and_unsupported_files_are_explicit(self):
        for name, raw, status in [("empty.txt", b"  \n", "failed"), ("broken.docx", b"bad", "failed"),
                                  ("broken.pdf", b"bad", "failed"), ("binary.txt", b"\xff", "failed"),
                                  ("empty.docx", docx_bytes(), "failed"), ("x.bin", b"data", "unsupported")]:
            with self.subTest(name=name):
                case = self.upload(name, raw)
                file = case["files"][0]
                self.assertEqual(file["parse_status"], status)
                self.assertIsNone(file["parsed_text_path"])
                self.assertTrue(file["error_message"])
                self.assertEqual(self.client.get(
                    f"/api/cases/{case['case_id']}/files/{file['file_id']}/text").status_code, 422)

    def test_missing_or_unreadable_cache_never_reparses_original(self):
        for value in (None, b"", b"\xff"):
            case = self.upload("sample.docx", docx_bytes("完整正文"))
            file = case["files"][0]
            target = self.data_dir / case["case_id"] / file["parsed_text_path"]
            if value is None:
                target.unlink()
            else:
                target.write_bytes(value)
            with patch.object(file_parsers, "parse_text", side_effect=AssertionError("unexpected parse")):
                response = self.client.get(f"/api/cases/{case['case_id']}/extraction-input")
                self.assertEqual(response.status_code, 409, response.text)
                self.assertIn(response.json()["detail"]["code"], {"parsed_cache_empty", "parsed_cache_unavailable"})

    def test_cached_input_is_read_repeatedly_without_parsing(self):
        with patch.object(file_parsers, "parse_text", wraps=file_parsers.parse_text) as parser:
            case = self.upload("sample.docx", docx_bytes("第一段", "第二段😀"))
            self.assertEqual(parser.call_count, 1)
            for _ in range(3):
                response = self.client.get(f"/api/cases/{case['case_id']}/extraction-input")
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(response.json()["files"][0]["text"], "第一段\n第二段😀")
            self.assertEqual(parser.call_count, 1)

    def test_cache_text_endpoint_returns_full_document(self):
        case = self.upload()
        response = self.client.get(f"/api/cases/{case['case_id']}/files/{case['files'][0]['file_id']}/text")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["text"], "甲方付款😀\n乙方收款")

    def test_docx_tables_remain_in_document_order(self):
        document = Document()
        document.add_paragraph("开始")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "金额"
        document.add_paragraph("结束")
        output = io.BytesIO()
        document.save(output)
        self.assertEqual(file_parsers.parse_text(FileType.docx, output.getvalue()), "开始\n项目\t金额\n结束")

    def test_pdf_text_and_page_markers_are_cached(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (Full document text) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = io.BytesIO()
        writer.write(output)
        case = self.upload("text.pdf", output.getvalue())
        self.assertEqual(case["files"][0]["parse_status"], "parsed")
        text = (self.data_dir / case["case_id"] / case["files"][0]["parsed_text_path"]).read_text(encoding="utf-8")
        self.assertTrue(text.startswith("[Page 1]\n"))
        self.assertIn("Full document text", text)

    def test_blank_pdf_is_not_success(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        output = io.BytesIO()
        writer.write(output)
        self.assertEqual(self.upload("blank.pdf", output.getvalue())["files"][0]["parse_status"], "failed")

    def test_bom_and_markdown_have_stable_utf8_cache(self):
        case = self.upload("notes.md", b"\xef\xbb\xbf# Heading\r\ntext")
        self.assertEqual(case["files"][0]["file_type"], "markdown")
        self.assertEqual(build_extraction_input(case_store.load_case(case["case_id"])).files[0].text, "# Heading\ntext")

    def test_rejects_nonmultipart_and_invalid_forms_without_writes(self):
        self.assertEqual(self.client.post("/api/cases", content=b"not-json",
                                          headers={"content-type": "text/plain"}).status_code, 415)
        self.assertEqual(self.client.post("/api/cases", content=b"{",
                                          headers={"content-type": "application/json"}).status_code, 422)
        for data, files in [({}, {"files": ("f.txt", b"text")}),
                             ({"case_title": "  "}, {"files": ("f.txt", b"text")}),
                             ({"case_title": "title"}, {"wrong": ("f.txt", b"text")})]:
            with self.subTest(data=data):
                self.assertEqual(self.client.post("/api/cases", data=data, files=files).status_code, 422)
        self.assertEqual(list(self.data_dir.iterdir()), [])

    def test_empty_case_create_and_rename_persist_without_changing_identity(self):
        response = self.client.post("/api/cases", json={"case_title": "未命名案件"})
        self.assertEqual(response.status_code, 201, response.text)
        created = response.json()["case"]
        self.assertEqual(created["case_title"], "未命名案件")
        self.assertEqual(created["files"], [])
        self.assertEqual(created["info_units"], [])
        self.assertEqual(self.client.get("/api/case").json()["case"], created)
        listing = self.client.get("/api/cases").json()
        self.assertEqual(listing["current_case_id"], created["case_id"])
        self.assertEqual([case["case_id"] for case in listing["cases"]], [created["case_id"]])

        renamed_response = self.client.patch(
            f"/api/cases/{created['case_id']}", json={"case_title": "租赁争议"})
        self.assertEqual(renamed_response.status_code, 200, renamed_response.text)
        renamed = renamed_response.json()["case"]
        self.assertEqual(renamed["case_id"], created["case_id"])
        self.assertEqual(renamed["created_at"], created["created_at"])
        self.assertEqual(renamed["case_title"], "租赁争议")
        self.assertEqual(renamed["files"], created["files"])
        self.assertEqual(self.client.get(f"/api/cases/{created['case_id']}").json()["case"], renamed)
        self.assertEqual(self.client.get("/api/cases").json()["cases"][0]["case_title"], "租赁争议")

        before_invalid = (self.data_dir / created["case_id"] / "case.json").read_bytes()
        invalid = self.client.patch(f"/api/cases/{created['case_id']}", json={"case_title": "  "})
        self.assertEqual(invalid.status_code, 422, invalid.text)
        self.assertEqual((self.data_dir / created["case_id"] / "case.json").read_bytes(), before_invalid)

    def test_remaining_unconnected_pipelines_return_501_and_never_write_results(self):
        case = self.upload()
        path = self.data_dir / case["case_id"] / "case.json"
        before = path.read_bytes()
        inputs = {
            "/api/reasoning/run": {"case_id": case["case_id"], "question": "争点", "selected_info_unit_ids": [], "selected_relation_ids": []},
            "/api/law/retrieve": {"case_id": case["case_id"], "issue": "争点", "selected_info_unit_ids": [], "max_results": 5},
        }
        for route, data in inputs.items():
            with self.subTest(route=route):
                response = self.client.post(route, json=data)
                self.assertEqual(response.status_code, 501, response.text)
                self.assertEqual(response.json()["detail"]["code"], "pipeline_not_implemented")
        self.assertEqual(path.read_bytes(), before)

    def test_invalid_selection_and_missing_case_are_errors(self):
        case = self.upload()
        stored = case_store.load_case(case["case_id"])
        for selected in ([], ["unknown"], [case["files"][0]["file_id"]] * 2):
            with self.subTest(selected=selected), self.assertRaises(AppError):
                build_extraction_input(stored, selected)
        self.assertEqual(self.client.get("/api/cases/missing").status_code, 404)

    def test_append_write_failure_rolls_back_artifacts_and_snapshot(self):
        case = self.upload()
        stored = case_store.load_case(case["case_id"])
        directory = self.data_dir / case["case_id"]
        before = {str(p.relative_to(directory)): p.read_bytes() for p in directory.rglob("*") if p.is_file()}
        real_write = case_service.atomic_write
        def failing_write(path, raw):
            if path.parent.name == "parsed":
                raise PermissionError("disk write denied")
            real_write(path, raw)
        with patch.object(case_service, "atomic_write", side_effect=failing_write):
            with self.assertRaises(PermissionError):
                asyncio.run(append_files_to_case(stored, [UploadFile(io.BytesIO(b"new"), filename="new.txt")], {"new.txt": "document"}))
        after = {str(p.relative_to(directory)): p.read_bytes() for p in directory.rglob("*") if p.is_file()}
        self.assertEqual(after, before)
        self.assertEqual(len(stored.files), 1)

    def test_upload_read_failure_does_not_leave_partial_case(self):
        async def unreadable():
            raise OSError("upload stream failed")
        broken = UploadFile(io.BytesIO(b"bad"), filename="bad.txt")
        broken.read = unreadable
        with self.assertRaises(OSError):
            asyncio.run(case_service.create_case_from_uploads("case", [
                UploadFile(io.BytesIO(b"text"), filename="first.txt"), broken], {"first.txt": "document", "bad.txt": "record"}))
        self.assertEqual(list(self.data_dir.iterdir()), [])

    def test_storage_failure_returns_500_without_fake_success(self):
        with patch.object(case_service, "atomic_write", side_effect=PermissionError("denied")):
            response = self.client.post("/api/cases", data={"case_title": "case", "source_kinds": '{"a.txt":"document"}'}, files={"files": ("a.txt", b"text")})
        self.assertEqual(response.status_code, 500, response.text)
        self.assertEqual(response.json()["detail"]["code"], "storage_error")

    def test_pointer_write_failure_rolls_back_case_snapshot_and_files(self):
        first = self.upload()
        before = {str(p.relative_to(self.data_dir)): p.read_bytes()
                  for p in self.data_dir.rglob("*") if p.is_file()}
        original_write = case_store.atomic_write
        def failing_pointer(path, raw):
            if path.name == 'current_case_id.txt':
                raise PermissionError('pointer write failed')
            original_write(path, raw)
        with patch.object(case_store, 'atomic_write', side_effect=failing_pointer):
            response = self.client.post(f"/api/cases/{first['case_id']}/files",
                                        data={"source_kinds": '{"new.txt":"document"}'},
                                        files={'files': ('new.txt', b'new content')})
        self.assertEqual(response.status_code, 500, response.text)
        after = {str(p.relative_to(self.data_dir)): p.read_bytes()
                 for p in self.data_dir.rglob('*') if p.is_file()}
        self.assertEqual(after, before)

    def test_paths_are_case_relative_and_traversal_is_rejected(self):
        case = self.upload("../../outside.txt", b"content")
        file = case["files"][0]
        self.assertTrue(file["original_path"].startswith("upload_files/file-"))
        self.assertFalse((self.data_dir.parent / "outside.txt").exists())
        for cid in ("..", "a/b", "C:\\outside", "/absolute"):
            with self.subTest(cid=cid), self.assertRaises(AppError):
                case_store.case_dir(cid)
        for path in ("../outside", "C:/outside", "parsed/../../outside", "parsed\\x.txt"):
            with self.subTest(path=path), self.assertRaises(AppError):
                case_store.artifact_path(case["case_id"], path)

    def test_invalid_snapshot_is_reported_without_silent_skip(self):
        directory = self.data_dir / "case-invalid"
        directory.mkdir()
        (directory / "case.json").write_text('{"unexpected": true}', encoding="utf-8")
        self.assertEqual(self.client.get("/api/cases/case-invalid").status_code, 409)
        self.assertEqual(self.client.get("/api/cases").status_code, 409)

    def test_case_listing_selection_workspace_and_delete_use_current_contract(self):
        first = self.upload()
        second = self.upload()
        self.assertEqual(self.client.get("/api/cases").json()["current_case_id"], second["case_id"])
        self.assertEqual(self.client.post(f"/api/cases/{first['case_id']}/select").status_code, 200)
        response = self.client.patch(f"/api/cases/{first['case_id']}/workspace", json=first["workspace_state"])
        self.assertEqual(response.status_code, 200, response.text)
        bad_state = {**first["workspace_state"], "selected_node_ids": ["unknown"]}
        self.assertEqual(self.client.patch(f"/api/cases/{first['case_id']}/workspace", json=bad_state).status_code, 422)
        self.assertEqual(self.client.delete(f"/api/cases/{first['case_id']}").status_code, 200)
        self.assertIsNone(self.client.get("/api/case").json()["case"])
        self.assertEqual(len(self.client.get("/api/cases").json()["cases"]), 1)


if __name__ == "__main__":
    unittest.main()
